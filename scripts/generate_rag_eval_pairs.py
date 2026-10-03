#!/usr/bin/env python3
"""
Generate candidate issue pairs for the GST NoticeDesk RAG Labelled Evaluation Set.
Populates NoticeDesk_RAG_Labelled_Set_Template.xlsx Data tab starting at row 3.
"""

import os
import json
import re
from collections import defaultdict
import docx
import openpyxl

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(WORKSPACE_DIR, 'NoticeDesk_RAG_Labelled_Set_Template.xlsx')
CORPUS_DIR = os.path.join(WORKSPACE_DIR, 'NoticeDesk_GST_Corpus_1300_Paired', 'NoticeDesk_GST_Corpus_1-300_Paired')
MANIFEST_PATH = os.path.join(CORPUS_DIR, 'pairing_manifest.jsonl')

def tokenize(text):
    stopwords = set([
        'the', 'a', 'an', 'and', 'or', 'in', 'on', 'at', 'to', 'for', 'of', 'with', 'by', 
        'is', 'are', 'was', 'were', 'be', 'been', 'being', 'that', 'this', 'it', 'from',
        'as', 'under', 'proposes', 'alleging', 'alleges', 'allegation', 'company', 'taxpayer',
        'notice', 'show', 'cause', 'form', 'dated', 'period', 'amount', 'specified', 'section',
        'gst', 'gstr', 'rule', 'rules', 'cgst', 'sgst', 'igst', 'act', '2017', 'ltd', 'pvt'
    ])
    words = re.findall(r'\b[a-z0-9]{3,}\b', text.lower())
    return set(w for w in words if w not in stopwords)

def jaccard_sim(set1, set2):
    if not set1 or not set2:
        return 0.0
    return len(set1 & set2) / len(set1 | set2)

def clean_pii(text):
    if not text:
        return ''
    t = text
    # Remove financial amounts & Rs/INR currency markers
    t = re.sub(r'\bapprox(imately)?\s*Rs\.?\s*[\d,]+', 'specified amount', t, flags=re.IGNORECASE)
    t = re.sub(r'\bRs\.?\s*[\d,]+(\.\d+)?', 'specified amount', t, flags=re.IGNORECASE)
    t = re.sub(r'\bINR\s*[\d,]+', 'specified amount', t, flags=re.IGNORECASE)
    t = re.sub(r'\b(Rs\.?|INR|₹)\b', 'specified amount', t, flags=re.IGNORECASE)
    t = re.sub(r'\b\d{2}[A-Z]{5}\d{4}[A-Z]{1}[A-Z0-9]{1}Z[A-Z0-9]{1}\b', '[GSTIN]', t)
    
    # Remove company names & M/s placeholders
    t = re.sub(r'M/s\s*(\[_*\]|_+|[A-Za-z0-9\.\s]+?(Private|Pvt|Limited|Ltd|Traders|Enterprises|Industries)?)', 'the supplier', t, flags=re.IGNORECASE)
    t = re.sub(r'Noticedesk\s+(Technologies\s+Private\s+Limited)?', 'Taxpayer', t, flags=re.IGNORECASE)
    
    # Remove identifiers & placeholders
    t = re.sub(r'GSTIN:?\s*\[.*?\]', '', t)
    t = re.sub(r'DIN:?\s*\[.*?\]', '', t)
    t = re.sub(r'\[_+\]', '', t)
    
    # Fix spacing
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def classify_topic(issue_text, serial, section):
    t_lower = issue_text.lower()
    
    # 1. Fake Invoice
    if 'fake' in t_lower or 'bogus' in t_lower or 'non-existent' in t_lower or '07a' in serial.lower() or '08' in serial.lower():
        return 'fake_invoice'
    # 2. Blocked Credit under Section 17(5) - Checked before Rule 86A
    elif '17(5)' in t_lower or 'blocked credit' in t_lower or 'plant and machinery' in t_lower or 'construction of immovable' in t_lower:
        return 'blocked_credit_17_5'
    # 3. Rule 86A Credit Blocking - Requires explicit 86a
    elif '86a' in t_lower:
        return 'rule_86a_credit_block'
    # 4. E-Way Bill / Goods Detention under Section 129 - Requires e-way/transit keywords + 129/detention context
    elif ('eway' in t_lower or 'e-way' in t_lower or 'mov-06' in t_lower or 'mov-07' in t_lower or 'transit' in t_lower or 'vehicle' in t_lower) and ('129' in section or 'deten' in t_lower or 'intercep' in t_lower or 'conveyance' in t_lower):
        return 'eway_detention'
    # 5. Interest Demand under Section 50 - Must be primary subject
    elif 'section 50' in t_lower or 'interest under' in t_lower or 'interest on' in t_lower or 'interest gross' in t_lower or 'net liability' in t_lower or 'interest delayed' in t_lower:
        return 'interest_demand'
    # 6. ITC Mismatch
    elif '2a' in t_lower or '2b' in t_lower or 'mismatch' in t_lower or 'reconciliation' in t_lower or 'rule 36(4)' in t_lower:
        return 'itc_mismatch'
    # 7. Cancellation / Revocation of Registration
    elif 'cancellation' in t_lower or 'revocation' in t_lower or 'reg17' in t_lower or 'reg-17' in t_lower or 'reg-21' in t_lower:
        return 'cancellation_registration'
    # 8. Refund / Drawback
    elif 'refund' in t_lower or 'drawback' in t_lower or 'section 54' in t_lower or 'rfd-01' in t_lower or 'rfd-08' in t_lower:
        return 'refund_claim'
    # 9. Wrong Head Tax Payment
    elif 'wrong head' in t_lower or 'section 77' in t_lower or ('cgst' in t_lower and 'igst' in t_lower and 'head' in t_lower):
        return 'wrong_head_tax'
    else:
        return 'general_tax_demand'

def load_corpus():
    with open(MANIFEST_PATH) as f:
        manifest = [json.loads(line) for line in f if line.strip()]

    docs = []
    for item in manifest:
        serial = item['serial']
        n_fn = item['notice_file']
        d_fn = item['draft_file']
        notice_kind = item['notice_kind']
        
        n_path = os.path.join(CORPUS_DIR, 'notices', n_fn)
        doc = docx.Document(n_path)
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        subj = ''
        stat_ref = ''
        allegation = ''
        for p in paras:
            if p.startswith('Subject:'):
                subj = p.replace('Subject:', '').strip()
            elif p.startswith('Statutory reference:'):
                stat_ref = p.replace('Statutory reference:', '').strip()
            elif p.startswith('ALLEGATIONS / GROUNDS') or p.startswith('Allegations:'):
                idx = paras.index(p)
                if idx + 1 < len(paras):
                    allegation = paras[idx + 1]
        
        if not allegation:
            d_path = os.path.join(CORPUS_DIR, 'drafts', d_fn)
            d_doc = docx.Document(d_path)
            d_paras = [p.text.strip() for p in d_doc.paragraphs if p.text.strip()]
            for p in d_paras:
                if p == 'Synopsis' or p.startswith('Subject:'):
                    idx = d_paras.index(p)
                    if idx + 1 < len(d_paras):
                        allegation = d_paras[idx + 1]
                    break

        title = clean_pii(subj)
        desc = clean_pii(allegation)
        issue_text = f"{title} — {desc}" if desc else title

        sec_match = re.search(r'Section\s+(\d+[A-Z]?(?:\(\d+\))?(?:\([a-z]+\))?)', subj + ' ' + stat_ref + ' ' + n_fn, re.IGNORECASE)
        section = sec_match.group(0) if sec_match else ('Section 73' if notice_kind == 'scn_73' else ('Section 74' if notice_kind == 'scn_74' else 'Other'))

        doc_entry = {
            'serial': serial,
            'notice_kind': notice_kind,
            'section': section,
            'title': title,
            'desc': desc,
            'issue_text': issue_text,
            'tokens': tokenize(issue_text),
            'source_ref': f"GST_CORPUS pair #{serial}"
        }
        doc_entry['topic'] = classify_topic(issue_text, serial, section)
        docs.append(doc_entry)
    
    return docs

def generate_pairs(docs):
    pairs = []
    seen_pairs = set()
    global_doc_counts = defaultdict(int)

    def add_pair(d1, d2, pair_type, max_global_reuse=4):
        pair_key = tuple(sorted([d1['serial'], d2['serial']]))
        if pair_key in seen_pairs:
            return False
        if global_doc_counts[d1['serial']] >= max_global_reuse or global_doc_counts[d2['serial']] >= max_global_reuse:
            return False
        seen_pairs.add(pair_key)
        global_doc_counts[d1['serial']] += 1
        global_doc_counts[d2['serial']] += 1
        pairs.append({
            'doc1': d1,
            'doc2': d2,
            'pair_type': pair_type
        })
        return True

    # 1. likely_same (~64 pairs)
    # Target: SAME section AND SAME specific issue topic (excluding generic fallbacks)
    topic_groups = defaultdict(list)
    for d in docs:
        if d['topic'] not in ('general_tax_demand', 'other'):
            topic_groups[(d['section'], d['topic'])].append(d)

    likely_same_count = 0
    for (sec, topic), grp in topic_groups.items():
        if len(grp) >= 2:
            for i in range(len(grp)):
                for j in range(i + 1, len(grp)):
                    if likely_same_count >= 64:
                        break
                    if add_pair(grp[i], grp[j], 'likely_same', max_global_reuse=3):
                        likely_same_count += 1
                if likely_same_count >= 64:
                    break

    # 2. hard_same_words_diff_law (~32 pairs)
    # Target: DIFFERENT section OR DIFFERENT notice_kind, AND (SAME topic OR Jaccard >= 0.25)
    t2_candidates = []
    for i in range(len(docs)):
        for j in range(i + 1, len(docs)):
            d1, d2 = docs[i], docs[j]
            diff_law = (d1['section'] != d2['section']) or (d1['notice_kind'] != d2['notice_kind'])
            if diff_law:
                sim = jaccard_sim(d1['tokens'], d2['tokens'])
                same_topic = (d1['topic'] == d2['topic']) and (d1['topic'] not in ('general_tax_demand', 'other'))
                if same_topic or sim >= 0.25:
                    t2_candidates.append((d1, d2, sim))

    t2_candidates.sort(key=lambda x: x[2], reverse=True)
    t2_count = 0
    for d1, d2, sim in t2_candidates:
        if t2_count >= 32:
            break
        if add_pair(d1, d2, 'hard_same_words_diff_law', max_global_reuse=3):
            t2_count += 1

    # 3. hard_diff_words_same_issue (~32 pairs) - STRICT SPECIFIC TOPIC FILTER
    # Target: SAME section AND SAME specific issue topic (EXCLUDING general_tax_demand/other fallback)
    t3_candidates = []
    for i in range(len(docs)):
        for j in range(i + 1, len(docs)):
            d1, d2 = docs[i], docs[j]
            same_specific_topic = (d1['section'] == d2['section']) and (d1['topic'] == d2['topic']) and (d1['topic'] not in ('general_tax_demand', 'other'))
            if same_specific_topic:
                sim = jaccard_sim(d1['tokens'], d2['tokens'])
                if 0.05 <= sim < 0.35:
                    t3_candidates.append((d1, d2, sim))

    t3_candidates.sort(key=lambda x: x[2])
    t3_count = 0
    for d1, d2, sim in t3_candidates:
        if t3_count >= 32:
            break
        if add_pair(d1, d2, 'hard_diff_words_same_issue', max_global_reuse=4):
            t3_count += 1

    # 4. random_control (~32 pairs)
    # Target: DIFFERENT sections AND DIFFERENT topics
    t4_count = 0
    import random
    random.seed(42)
    doc_indices = list(range(len(docs)))
    random.shuffle(doc_indices)
    
    for i in range(len(doc_indices)):
        for j in range(i + 1, len(doc_indices)):
            if t4_count >= 32:
                break
            d1 = docs[doc_indices[i]]
            d2 = docs[doc_indices[j]]
            if d1['section'] != d2['section'] and d1['topic'] != d2['topic']:
                if add_pair(d1, d2, 'random_control', max_global_reuse=4):
                    t4_count += 1
        if t4_count >= 32:
            break

    return pairs

def main():
    print(f"Loading corpus from {CORPUS_DIR}...")
    docs = load_corpus()
    print(f"Loaded {len(docs)} documents.")

    pairs = generate_pairs(docs)
    print(f"Generated {len(pairs)} candidate pairs total.")

    type_counts = defaultdict(int)
    for p in pairs:
        type_counts[p['pair_type']] += 1
    
    print("Breakdown by pair type:")
    for ptype, cnt in sorted(type_counts.items()):
        print(f"  - {ptype}: {cnt}")

    # Confidentiality Audit Scan
    print("\nRunning Confidentiality Audit Scan over generated pairs...")
    pii_violations = []
    for i, p in enumerate(pairs, start=1):
        pair_id = f"P{i:03d}"
        t1, t2 = p['doc1']['issue_text'], p['doc2']['issue_text']
        for text, field in [(t1, 'Issue 1 Text'), (t2, 'Issue 2 Text')]:
            if re.search(r'Rs\.?\s*\d|₹\s*\d|\b\d{5,}\b', text):
                pii_violations.append((pair_id, field, 'amount', text))
            if re.search(r'M/s\b|Noticedesk|Private Limited', text, re.IGNORECASE):
                pii_violations.append((pair_id, field, 'company_name', text))
            if re.search(r'\b\d{2}[A-Z]{5}\d{4}[A-Z]{1}[A-Z0-9]{1}Z[A-Z0-9]{1}\b', text):
                pii_violations.append((pair_id, field, 'gstin', text))

    if pii_violations:
        print(f"WARNING: Found {len(pii_violations)} PII violations:")
        for v in pii_violations:
            print(f"  - [{v[0]}] {v[1]} ({v[2]}): {v[3][:100]}")
    else:
        print("CONFIDENTIALITY AUDIT PASSED: 0 PII, amount, or GSTIN violations detected.")

    # Write to Excel Template
    print(f"\nWriting to Excel template: {TEMPLATE_PATH}...")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb['Data']

    for idx, p in enumerate(pairs, start=1):
        row = idx + 2 # row 3 is first pair
        pair_id = f"P{idx:03d}"
        d1, d2 = p['doc1'], p['doc2']
        ptype = p['pair_type']

        ws.cell(row=row, column=1, value=pair_id)
        ws.cell(row=row, column=2, value=d1['issue_text'])
        ws.cell(row=row, column=3, value=d1['source_ref'])
        ws.cell(row=row, column=4, value=d2['issue_text'])
        ws.cell(row=row, column=5, value=d2['source_ref'])
        ws.cell(row=row, column=6, value=ptype)

    wb.save(TEMPLATE_PATH)
    print(f"Successfully saved {len(pairs)} rows to {TEMPLATE_PATH}.")

if __name__ == '__main__':
    main()
